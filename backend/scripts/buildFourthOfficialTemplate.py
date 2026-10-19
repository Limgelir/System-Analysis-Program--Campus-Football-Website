"""Rebuild the editable Word template from the supplied Excel form.

Development-only: needs openpyxl and python-docx. The server uses the committed
DOCX and does not need Python, Excel, or Microsoft Word.
"""
from io import BytesIO
from pathlib import Path

import openpyxl
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / '四官单模板' / '四官单模板--11人制版.xlsx'
DEST = ROOT / 'backend' / 'templates' / 'fourth-official-11.docx'
sheet = openpyxl.load_workbook(SOURCE).worksheets[0]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = Mm(10)
section.left_margin = section.right_margin = Mm(10)
section.header_distance = section.footer_distance = Mm(0)
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(9)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
style.paragraph_format.line_spacing = 1
doc.core_properties.title = '11人制四官单'
doc.core_properties.author = 'TJFA'

fields = {'U2': 'field_order', 'C3': 'date', 'I3': 'time', 'N3': 'supervisor',
          'S3': 'venue', 'E4': 'main', 'L4': 'assistant1', 'P4': 'assistant2',
          'W4': 'fourth', 'A5': 'team_a', 'S5': 'team_b',
          'I5': 'a_shirt', 'I6': 'a_shorts', 'I7': 'a_socks',
          'Q5': 'b_shirt', 'Q6': 'b_shorts', 'Q7': 'b_socks'}
for coord, field in list(fields.items()):
    cell = sheet[coord]
    fields[sheet.cell(cell.row + 39, cell.column).coordinate] = field
for i in range(30):
    row = 9 + i
    for col, name in [('J', 'a_name'), ('L', 'a_number'), ('H', 'a_suspended'),
                      ('P', 'b_name'), ('O', 'b_number'), ('R', 'b_suspended')]:
        fields[f'{col}{row}'] = f'{name}_{i + 1}'

def element(name, **attrs):
    node = OxmlElement('w:' + name)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node

widths = [sheet.column_dimensions[openpyxl.utils.get_column_letter(c)].width * 7 + 5 for c in range(1, 27)]
total_width = int(Mm(190).twips)
widths = [round(w / sum(widths) * total_width) for w in widths]
widths[-1] += total_width - sum(widths)
merged = list(sheet.merged_cells.ranges)
logo = sheet._images[0]._data()

def add_page(start, end):
    table_xml = element('tbl')
    props = element('tblPr')
    props.append(element('tblW', w=total_width, type='dxa'))
    props.append(element('tblLayout', type='fixed'))
    borders = element('tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        borders.append(element(side, val='single', sz=4, color='000000'))
    props.append(borders)
    margins = element('tblCellMar')
    for side in ['top', 'bottom', 'left', 'right']:
        margins.append(element(side, w=0 if side in ['top', 'bottom'] else 12, type='dxa'))
    props.append(margins)
    table_xml.append(props)
    grid = element('tblGrid')
    for w in widths:
        grid.append(element('gridCol', w=w))
    table_xml.append(grid)
    doc._body._body.insert_element_before(table_xml, 'w:sectPr')
    table = Table(table_xml, doc._body)
    for row_no in range(start, end + 1):
        tr = element('tr')
        pr = element('trPr')
        pr.append(element('cantSplit'))
        pr.append(element('trHeight', val=round((sheet.row_dimensions[row_no].height or 15) * .92 * 20), hRule='exact'))
        tr.append(pr)
        table_xml.append(tr)
        col = 1
        while col <= 26:
            merge = next((m for m in merged if m.min_row <= row_no <= m.max_row and m.min_col <= col <= m.max_col), None)
            span = merge.max_col - col + 1 if merge else 1
            cell_xml = element('tc')
            cell_pr = element('tcPr')
            cell_pr.append(element('tcW', w=sum(widths[col - 1:col - 1 + span]), type='dxa'))
            if span > 1:
                cell_pr.append(element('gridSpan', val=span))
            if merge and merge.max_row > merge.min_row:
                cell_pr.append(element('vMerge', val='restart' if row_no == merge.min_row else 'continue'))
            cell_pr.append(element('vAlign', val='center'))
            cell_xml.append(cell_pr)
            cell_xml.append(element('p'))
            tr.append(cell_xml)
            from docx.table import _Cell
            cell = _Cell(cell_xml, table)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.widow_control = False
            source = sheet.cell(row_no, col)
            continuation = merge and row_no > merge.min_row
            if not continuation:
                if source.coordinate in ('A1', 'A40'):
                    picture = p.add_run().add_picture(BytesIO(logo), width=Pt(sum(widths[:8]) / 20 - 5), height=Pt(58))
                    # Preserve the Excel drawing's existing crop; bitmap is unchanged.
                    fill = picture._inline.graphic.graphicData.pic.blipFill
                    crop = OxmlElement('a:srcRect')
                    for key, value in {'l': '9390', 't': '23695', 'r': '5164', 'b': '24173'}.items():
                        crop.set(key, value)
                    fill.insert(1, crop)
                    picture._inline.docPr.set('descr', '同济大学足球协会 TJFA')
                else:
                    value = '{{' + fields[source.coordinate] + '}}' if source.coordinate in fields else str(source.value or '').strip()
                    if source.coordinate == 'I40':
                        value = sheet['I1'].value
                    run = p.add_run(value)
                    run.font.name = '宋体'
                    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
                    run.font.size = Pt(15 if source.coordinate in ('I1', 'I40') else 9)
                    run.bold = bool(source.font.b) if source.coordinate not in fields else False
            col += span

add_page(1, 39)
gap = doc.add_paragraph()
gap.paragraph_format.space_before = gap.paragraph_format.space_after = Pt(0)
gap.paragraph_format.line_spacing = Pt(1)
gap.add_run().add_break(__import__('docx.enum.text', fromlist=['WD_BREAK']).WD_BREAK.PAGE)
add_page(40, 69)
tail = doc.add_paragraph()
tail.paragraph_format.line_spacing = Pt(1)
DEST.parent.mkdir(parents=True, exist_ok=True)
doc.save(DEST)
print(DEST)
